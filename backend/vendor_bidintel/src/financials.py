"""Financial statements (7) and annual reports (2) -> the `financials` and
`annual_reports` fact blocks.

BRIEFING.md s1 warns that "financial statements change format across eras" and
that "a parser tuned to the newest statements will misread the oldest". All 7
statements were read in full before this file was written. Three layout eras
exist (see `detect_era`), and the one that matters is the middle era: it renders
the tables in a narrower column, so labels wrap and the wrapped fragment lands
*after* the values once blocks are sorted by position::

    era 1 / 3                        era 2
    Total Revenue from Operations (A)    Total Revenue from Operations
    33,649                               33,649
    28,585                               28,585
    Tax Expense (current + deferred)     (A)
    766                                  ...
                                         Tax Expense (current +
                                         766
                                         deferred)

So the parser never matches a whole label. It matches a label that is either the
whole line or the head of the line up to an opening parenthesis, and then reads
the run of numeric lines that follows. That rule is era-independent: no branch in
this module keys on a document id, a year, or an era. `detect_era` exists only so
the audit and the tests can prove every era is covered.

Two facts about these documents drive the field choices, and both were
established against independent renderings rather than assumed:

* **`contract_revenue` is the turnover figure, not `revenue_operations`.** The
  statements print `Contract Revenue (EPC)`, an `Other Operating Revenue` line
  that is exactly 3% of it, and their sum as "Total Revenue from Operations (A)".
  The annual reports' seven-year summary and the trial-balance workbook both give
  net revenue from operations in raw rupees, and both equal the *contract
  revenue* line, in 7 of 7 years, to the rupee. Nothing anywhere corroborates the
  3% uplift.
* **`pbt` is recorded as printed, not as A - B.** The header says
  "C. Profit Before Tax (A - B)" but the printed figure never equals A - B, in
  any of the 7. The printed PBT/tax/PAT triple is internally consistent, drives
  the reserves movement, and its PAT matches the annual reports and the trial
  balance to the rupee in 7 of 7 years. A - B matches nothing and is a fixed
  19.42% of revenue every year. See `docs/experiment_log.md`.

The statements' *balance sheet* extract reconciles with nothing (0 of 7 against
the trial balance) and its "previous year" column is a fabricated 0.85x of the
current column rather than the prior year's figures. Both are transcribed
faithfully, and both are flagged on every audit run.
"""
from __future__ import annotations

import json
import pathlib
import re
import sys

REPO = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))

import money  # noqa: E402
import pdftext  # noqa: E402

OUT = REPO / "results" / "raw" / "financials.json"
WORKBOOKS = REPO / "dataset" / "documents" / "workbooks"

# A cell holding only a number: "4,817", "-2,803", "23356.37", "(1,064)".
_NUMERIC_LINE = re.compile(r"^\(?-?\s*\d[\d,]*(?:\.\d+)?\)?$")
# The first signed number in a cell, with the sign wherever it sits relative to
# any currency prefix: "Rs. -3,90,26,159" and "-2,803" are both negative.
_SIGNED = re.compile(r"(-)?\s*(\d[\d,]*(?:\.\d+)?)")
_UNIT_IN_TEXT = re.compile(r"\b(Cr(?:ore)?s?|Lakhs?|Lacs?)\b", re.IGNORECASE)
# The label separator. PyMuPDF renders these documents' en-dash as U+FFFD
# because the glyph is not mapped in the embedded font, so U+FFFD is a real
# separator here, not corruption to repair.
_SEPARATOR = re.compile(r"\s*[�–—]\s*|\s+-\s+")

# "Currency: INR (All amounts in Lakhs unless stated)". Read per document: the
# multiplier is never hardcoded, and "unless stated" is honoured by letting any
# unit written on an individual figure override this one.
_CURRENCY_LINE = re.compile(
    r"All amounts in\s+(?P<unit>[A-Za-z]+)\b", re.IGNORECASE)
_YEAR_ENDED = re.compile(
    r"year ended\s+\d{1,2}(?:st|nd|rd|th)?\s+March\s+(?P<year>\d{4})",
    re.IGNORECASE)
_FY_TOKEN = re.compile(r"\(?FY\s*(?P<start>\d{4})[�–—-](?P<end>\d{2})\)?")
# Every gap is \s+, never a literal space: the note wraps at a different word in
# each era ("...the order book of\napproximately 120 Cr" in eras 1 and 3,
# "...including the\norder book of approximately 817 Cr" in era 2), so a regex
# written with literal spaces silently finds the figure in only 3 of the 7.
_ORDER_BOOK = re.compile(
    r"order\s+book\s+of\s+approximately\s+(?P<amount>[\d,]+(?:\.\d+)?)"
    r"\s*(?P<unit>Cr(?:ore)?s?|Lakhs?|Lacs?)?\b", re.IGNORECASE)

# label -> field. The label is matched against a whole line, or against the head
# of a line up to " (" — never as a bare prefix, because decision_log.md s2
# records that bare prefix matching silently returns the wrong value.
_PNL_ROWS = [
    ("Contract Revenue", "contract_revenue"),
    ("Other Operating Revenue", "other_operating_revenue"),
    ("Total Revenue from Operations", "revenue_operations"),
    ("Cost of Materials Consumed", "cost_of_materials"),
    ("Sub-contracting & Labour", "subcontracting_labour"),
    ("Employee Benefit Expenses", "employee_benefit_expense"),
    ("Depreciation & Amortisation", "depreciation"),
    ("Other Expenses", "other_expenses"),
    ("Total Expenses", "total_expenses"),
]
# Section C prints a current-year column only — there is no comparative.
_PROFIT_ROWS = [
    ("Profit Before Tax", "pbt"),
    ("Tax Expense", "tax"),
    ("Profit After Tax", "pat"),
]
_BALANCE_ROWS = [
    ("Paid-up Capital", "paid_up_capital"),
    ("Reserves & Surplus", "reserves"),
    ("Long-term Borrowings", "long_term_borrowings"),
    ("Trade Payables", "trade_payables"),
    ("Other Current Liabilities", "other_current_liabilities"),
    ("Property, Plant & Equipment", "ppe"),
    ("Intangible Assets", "intangible_assets"),
    ("Inventories", "inventories"),
    ("Trade Receivables", "trade_receivables"),
    ("Cash & Bank Balances", "cash"),
]

CURRENT_ASSET_FIELDS = ("inventories", "trade_receivables", "cash")
CURRENT_LIABILITY_FIELDS = ("trade_payables", "other_current_liabilities")

# Every era marker below is literal text present in the document. Detection is
# never by document id or by year.
_ERAS = (
    ("allcaps-wrapped", ("A. REVENUE FROM OPERATIONS", "PARTICULARS")),
    ("capsheading-inline", ("A. REVENUE FROM OPERATIONS",)),
    ("titlecase-inline", ("A. Revenue from Operations",)),
)


class StatementError(ValueError):
    """A statement did not parse the way every read statement is shaped."""


# --------------------------------------------------------------------------- #
# primitives
# --------------------------------------------------------------------------- #

def detect_era(text: str) -> str:
    """Name the layout era from markers present in the text.

    Used by the audit and the tests to prove coverage; the parser itself is
    era-independent and never branches on the result.
    """
    for name, markers in _ERAS:
        if all(marker in text for marker in markers):
            return name
    return "unknown"


def document_unit(text: str) -> str | None:
    """The unit named on the document's own currency line, or None.

    None means figures are raw rupees unless a figure states its own unit —
    which is how the annual reports are written.
    """
    m = _CURRENCY_LINE.search(text)
    return m.group("unit") if m else None


def to_rupees(cell: str, default_unit: str | None) -> int | None:
    """One cell -> exact integer rupees, signed.

    A unit written on the cell wins over `default_unit`: that is what
    "All amounts in Lakhs *unless stated*" means, and Note 3's "order book of
    approximately 120 Cr" is the case it exists for. Scaling goes through
    `money.parse_money`, so the arithmetic is Decimal and never float.
    """
    if cell is None:
        return None
    m = _SIGNED.search(cell)
    if not m:
        return None
    negative = m.group(1) == "-" or (cell.strip().startswith("(")
                                     and cell.strip().endswith(")"))
    stated = _UNIT_IN_TEXT.search(cell)
    unit = stated.group(1) if stated else default_unit
    value = money.parse_money(f"{m.group(2)} {unit}" if unit else m.group(2))
    if value is None:
        return None
    return -value if negative else value


def _label_of(line: str) -> str:
    """The label a line carries, folded for comparison.

    Balance-sheet rows prefix the field with its statement grouping
    ("Non-Current Liabilities <sep> Long-term Borrowings"), so the label is the
    last separated segment.
    """
    return _SEPARATOR.split(line.strip())[-1].strip().lower()


def _matches(line: str, label: str) -> bool:
    """Whole-label match, tolerating a trailing parenthetical.

    Accepts "Contract Revenue (EPC)", "Total Expenses (B)" and the era-2
    fragment "Tax Expense (current +" for the label they belong to. Rejects the
    decision_log.md s2 failure mode where a label matches a longer field name:
    "Work" does not match "Work Category".
    """
    got, want = _label_of(line), label.lower()
    return got == want or got.startswith(want + " (")


def _row(lines: list[str], label: str, count: int, doc_id: str) -> list[str]:
    """The run of numeric cells that follows the line carrying `label`.

    Reading forward until the first non-numeric line is what makes the parser
    era-independent: in era 2 the orphaned wrap fragments ("(A)", "deferred)")
    are non-numeric, so they terminate the run instead of corrupting it.
    """
    hits = [i for i, line in enumerate(lines) if _matches(line, label)]
    if not hits:
        return []
    if len(hits) > 1:
        raise StatementError(
            f"{doc_id}: label {label!r} matched {len(hits)} lines; the parser "
            f"cannot tell which row is meant: "
            f"{[lines[i] for i in hits]}")
    cells: list[str] = []
    for line in lines[hits[0] + 1:]:
        if not _NUMERIC_LINE.match(line.strip()):
            break
        cells.append(line.strip())
        if len(cells) == count:
            break
    return cells


def _section(lines: list[str], start: str, end: str) -> list[str]:
    """The slice of `lines` between two heading markers present in the document.

    The annual reports repeat "Gross billings" and "Profit for the year" as a
    highlights row, as a seven-year-summary column header and as a cash-flow
    line, so a whole-document label search is ambiguous — `_row` refuses it
    rather than picking one. Bounding the search by the report's own headings
    resolves it without keying on a document id.
    """
    try:
        first = next(i for i, ln in enumerate(lines) if start in ln)
    except StopIteration:
        return []
    last = next((i for i, ln in enumerate(lines[first + 1:], first + 1)
                 if end in ln), len(lines))
    return lines[first:last]


def fiscal_year(text: str, doc_id: str) -> str:
    """Normalise to the schema's "2019-20", cross-checking three renderings.

    The statement says "for the financial year ended 31st March 2020", prints
    "(FY2019-20)" over the current-year column, and is filed as DOC-FS-2019. All
    three are checked against each other rather than one being assumed, because
    the task brief never promises the doc-id year is the FY start year.
    """
    m = _YEAR_ENDED.search(text)
    if not m:
        raise StatementError(f"{doc_id}: no 'year ended <date> March <year>' line")
    end = int(m.group("year"))
    fy = f"{end - 1}-{str(end)[2:]}"

    header = _FY_TOKEN.search(text)
    if header:
        printed = f"{header.group('start')}-{header.group('end')}"
        if printed != fy:
            raise StatementError(
                f"{doc_id}: column header says FY{printed} but the statement "
                f"covers the year ended March {end} (FY{fy})")

    stem = re.search(r"(\d{4})$", doc_id)
    if stem and int(stem.group(1)) != end - 1:
        raise StatementError(
            f"{doc_id}: doc-id year {stem.group(1)} is not the FY start year "
            f"{end - 1}; the id/FY correspondence assumed by the audit is broken")
    return fy


# --------------------------------------------------------------------------- #
# financial statements
# --------------------------------------------------------------------------- #

def parse_statement(doc_id: str, text: str) -> dict:
    """One financial statement -> one `financials` record."""
    lines = [ln for ln in text.replace("\f", "\n").splitlines() if ln.strip()]
    unit = document_unit(text)
    if unit is None:
        raise StatementError(
            f"{doc_id}: no 'All amounts in <unit>' currency line. Refusing to "
            f"guess a multiplier — a wrong one is off by 100x silently.")

    rec: dict = {"fy": fiscal_year(text, doc_id)}

    for label, field in _PNL_ROWS + _BALANCE_ROWS:
        cells = _row(lines, label, 2, doc_id)
        rec[field] = to_rupees(cells[0], unit) if len(cells) > 0 else None
        rec[field + "_previous"] = to_rupees(cells[1], unit) if len(cells) > 1 else None

    # Section C carries no comparative column, so there is no mirror to fill.
    # Leaving them absent rather than zero is the schema's rule: absence is a
    # real answer (BRIEFING.md s5).
    for label, field in _PROFIT_ROWS:
        cells = _row(lines, label, 1, doc_id)
        rec[field] = to_rupees(cells[0], unit) if cells else None

    # Note 3 states the order book in crore inside a lakh-denominated document.
    # Passing the unit the note itself writes is what "unless stated" requires;
    # falling back to the document unit would be wrong by 100x.
    m = _ORDER_BOOK.search(text)
    rec["order_book"] = (to_rupees(f"{m.group('amount')} {m.group('unit') or ''}",
                                   unit) if m else None)

    rec["source_doc"] = doc_id
    return rec


def derived(rec: dict) -> dict:
    """Ratios a bid desk asks for, computed from the record's own inputs.

    Deliberately *not* stored in the fact block — the router computes these. This
    exists so the audit can display the trend a human needs to eyeball, and so
    the formulas the router must implement are written down exactly once.
    """
    cap, res = rec.get("paid_up_capital"), rec.get("reserves")
    net_worth = None if cap is None or res is None else cap + res

    ca = [rec.get(f) for f in CURRENT_ASSET_FIELDS]
    cl = [rec.get(f) for f in CURRENT_LIABILITY_FIELDS]
    current_assets = None if any(v is None for v in ca) else sum(ca)
    current_liabilities = None if any(v is None for v in cl) else sum(cl)
    working_capital = (None if current_assets is None or current_liabilities is None
                       else current_assets - current_liabilities)
    current_ratio = (None if not current_liabilities
                     else current_assets / current_liabilities)
    return {
        "net_worth": net_worth,
        "current_assets": current_assets,
        "current_liabilities": current_liabilities,
        "working_capital": working_capital,
        "current_ratio": current_ratio,
    }


# --------------------------------------------------------------------------- #
# annual reports
# --------------------------------------------------------------------------- #

_AR_HIGHLIGHT_ROWS = [
    ("Gross billings", "gross_billings"),
    ("Net revenue from operations", "net_revenue"),
    ("Profit for the year", "profit"),
]
_AR_SUMMARY_COLUMNS = ("gross_billings", "net_revenue", "profit")
_FY_ROW = re.compile(r"^(?P<start>\d{4})[�–—-](?P<end>\d{2})$")


def seven_year_summary(text: str) -> dict[str, dict[str, int]]:
    """The annual reports' SEVEN-YEAR FINANCIAL SUMMARY, in raw rupees.

    An independent rendering of the same quantities the statements print in
    lakhs. Used to cross-check the statements and, within the report itself, to
    supply the finest rendering of the report's own year (decision_log.md s3).
    """
    lines = [ln.strip() for ln in text.replace("\f", "\n").splitlines() if ln.strip()]
    out: dict[str, dict[str, int]] = {}
    for i, line in enumerate(lines):
        m = _FY_ROW.match(line)
        if not m:
            continue
        cells = lines[i + 1:i + 4]
        if len(cells) < 3 or not all(_NUMERIC_LINE.match(c) for c in cells):
            continue
        # The summary rows are the only FY-keyed rows followed by three raw
        # numbers and a percentage; the segment tables are keyed by name.
        if i + 4 >= len(lines) or "%" not in lines[i + 4]:
            continue
        fy = f"{m.group('start')}-{m.group('end')}"
        out[fy] = {field: to_rupees(cell, None)
                   for field, cell in zip(_AR_SUMMARY_COLUMNS, cells)}
    return out


def parse_annual_report(doc_id: str, text: str) -> dict:
    """One annual report -> one `annual_reports` record.

    The report states its own year's figures twice: lakh-rounded in FINANCIAL
    HIGHLIGHTS and to the rupee in the seven-year summary. decision_log.md s3
    already settled what to do when one document renders a value twice — keep
    the finer one — so the summary wins where it covers the year.
    """
    lines = [ln for ln in text.replace("\f", "\n").splitlines() if ln.strip()]
    m = _YEAR_ENDED.search(text)
    if not m:
        raise StatementError(f"{doc_id}: no 'year ended <date> March <year>' line")
    end = int(m.group("year"))
    fy = f"{end - 1}-{str(end)[2:]}"

    # The annual reports have no document-wide currency line: each figure states
    # its own unit, or is raw rupees. Passing None is what keeps
    # "Rs. -3,90,26,159" from being multiplied by a lakh it never claimed.
    rec: dict = {"fy": fy}
    highlights = _section(lines, "FINANCIAL HIGHLIGHTS", "MANAGEMENT")
    if not highlights:
        raise StatementError(f"{doc_id}: no FINANCIAL HIGHLIGHTS section")
    for label, field in _AR_HIGHLIGHT_ROWS:
        cells = _row(highlights, label, 2, doc_id)
        rec[field] = to_rupees(cells[0], None) if cells else None

    for field, value in seven_year_summary(text).get(fy, {}).items():
        if value is not None:
            rec[field] = value

    rec["source_doc"] = doc_id
    return rec


# --------------------------------------------------------------------------- #
# module contract
# --------------------------------------------------------------------------- #

def build(corpus: dict[str, str]) -> dict[str, list[dict]]:
    """The `financials` and `annual_reports` blocks, in fiscal-year order."""
    financials = [parse_statement(doc_id, text) for doc_id, text
                  in pdftext.docs_of_type(corpus, "financial_statement").items()]
    reports = [parse_annual_report(doc_id, text) for doc_id, text
               in pdftext.docs_of_type(corpus, "annual_report").items()]
    return {
        "financials": sorted(financials, key=lambda r: r["fy"]),
        "annual_reports": sorted(reports, key=lambda r: r["fy"]),
    }


# --------------------------------------------------------------------------- #
# cross-checks used by the audit
# --------------------------------------------------------------------------- #

_TB_EXPENSE_ACCOUNTS = frozenset({
    "Depreciation Expense", "Equipment Hire Charges", "Materials Consumed",
    "Salaries & Site Overheads", "Site Services", "Specialist Works",
    "Subcontracting Charges",
})


def trial_balance_by_fy() -> dict[str, dict[str, int]]:
    """Per-FY revenue, expenses and profit rebuilt from the trial-balance
    workbook, in raw rupees. An independent rendering: a mismatch means one of
    the two parsers is wrong.

    Returns {} if the workbook or openpyxl is unavailable, so the audit degrades
    to reporting fewer checks rather than failing.
    """
    path = WORKBOOKS / "Trial_Balance_by_Year.xlsx"
    if not path.exists():
        return {}
    try:
        import openpyxl
    except ImportError:
        return {}

    wb = openpyxl.load_workbook(path, data_only=True)
    out: dict[str, dict[str, int]] = {}
    for sheet in wb.sheetnames:
        m = re.fullmatch(r"TB\s+(\d{4}-\d{2})", sheet.strip())
        if not m:
            continue
        revenue = expenses = 0
        balances: dict[str, int] = {}
        for row in wb[sheet].iter_rows(values_only=True):
            account, _debit, _credit, balance = (list(row) + [None] * 4)[:4]
            if not isinstance(account, str) or account in ("Account", "TOTAL"):
                continue
            if not isinstance(balance, (int, float)):
                continue
            account = account.replace("�", "-")
            balances[account] = int(balance)
            if account.startswith("Contract Revenue"):
                revenue -= int(balance)          # revenue sits as a credit
            elif account in _TB_EXPENSE_ACCOUNTS:
                expenses += int(balance)
        out[m.group(1)] = {"revenue": revenue, "expenses": expenses,
                           "profit": revenue - expenses, "balances": balances}
    return out


def _rounds_to_lakh(exact: int | None, lakh_rounded: int | None) -> bool:
    """Does an exact rupee figure round to the lakh figure the statement prints?"""
    if exact is None or lakh_rounded is None:
        return False
    return round(exact / 100_000) * 100_000 == lakh_rounded


def audit(blocks: dict[str, list[dict]]) -> None:
    """Coverage, the per-FY trend a human should eyeball, and every cross-check.

    Everything printed here is a check that can fail. The three that are *known*
    to fail — the balance sheet against the trial balance, the comparative
    column against the prior statement, and PBT against A - B — are printed as
    findings rather than suppressed, because they are the module's real
    limitations and they bear on what a router may safely compute.
    """
    financials = blocks.get("financials", [])
    reports = blocks.get("annual_reports", [])

    expected_fs = len(list((REPO / "dataset" / "documents" / "financial_statement")
                           .glob("*.pdf"))) or 7
    expected_ar = len(list((REPO / "dataset" / "documents" / "annual_report")
                           .glob("*.pdf"))) or 2

    print("=" * 78)
    print("financials.py audit")
    print("=" * 78)
    for name, records, expected in (("financials", financials, expected_fs),
                                    ("annual_reports", reports, expected_ar)):
        flag = "" if len(records) == expected else "   <-- INCOMPLETE"
        print(f"{name:<16} {len(records):>3} / {expected} records{flag}")

    # audit() takes only the blocks, per the module contract, but the era census
    # and the seven-year cross-check need the source text. Re-reading the cached
    # corpus is cheap; failing to is not a reason to abort build_facts.py's whole
    # audit pass, so the checks that need it degrade to a flagged line instead.
    try:
        corpus = pdftext.load_corpus()
    except Exception as exc:                                  # noqa: BLE001
        corpus = {}
        print(f"\ncorpus text unavailable ({exc}); era census and cross-check 1 "
              f"skipped   <-- INCOMPLETE")

    print("\nlayout eras (detected by marker text, never by doc id or year)")
    corpus_fs = pdftext.docs_of_type(corpus, "financial_statement")
    census: dict[str, list[str]] = {}
    for doc_id, text in corpus_fs.items():
        census.setdefault(detect_era(text), []).append(doc_id)
    for era, docs in sorted(census.items()):
        flag = "   <-- INCOMPLETE" if era == "unknown" else ""
        print(f"  {era:<20} {len(docs)} doc(s)  {', '.join(docs)}{flag}")

    print("\nfield population (financials)")
    fields = [k for k in (financials[0] if financials else {}) if k != "fy"]
    for field in fields:
        filled = sum(1 for r in financials if r.get(field) is not None)
        flag = "   <-- INCOMPLETE" if filled < len(financials) else ""
        print(f"  {field:<34} {filled}/{len(financials)}{flag}")

    print("\nper-FY sanity display (Rs crore; net worth = capital + reserves,")
    print("working capital = current assets - current liabilities)")
    head = ("fy", "turnover", "rev(A)", "PBT", "PAT", "net worth", "work.cap",
            "cur.ratio", "order bk")
    print("  " + "".join(f"{h:>11}" for h in head))
    for r in financials:
        d = derived(r)
        def cr(v):
            return "        n/a" if v is None else f"{v / 1e7:11.2f}"
        print("  " + f"{r['fy']:>11}" + cr(r["contract_revenue"])
              + cr(r["revenue_operations"]) + cr(r["pbt"]) + cr(r["pat"])
              + cr(d["net_worth"]) + cr(d["working_capital"])
              + (f"{d['current_ratio']:11.2f}" if d["current_ratio"] else "        n/a")
              + cr(r["order_book"]))
    negatives = [r["fy"] for r in financials if (derived(r)["net_worth"] or 0) < 0]
    if negatives:
        print(f"  NOTE: net worth is negative in {len(negatives)} of "
              f"{len(financials)} FYs: {', '.join(negatives)}. Every statement's "
              f"Note 3 nevertheless asserts 'positive net worth'.")

    print("\ncross-check 1: statement vs annual-report seven-year summary")
    corpus_ar = pdftext.docs_of_type(corpus, "annual_report")
    summary: dict[str, dict[str, int]] = {}
    for text in corpus_ar.values():
        summary.update(seven_year_summary(text))
    hits = {"contract_revenue": 0, "pat": 0, "revenue_operations": 0}
    for r in financials:
        row = summary.get(r["fy"])
        if not row:
            continue
        hits["contract_revenue"] += _rounds_to_lakh(row["net_revenue"],
                                                    r["contract_revenue"])
        hits["revenue_operations"] += _rounds_to_lakh(row["net_revenue"],
                                                      r["revenue_operations"])
        hits["pat"] += _rounds_to_lakh(row["profit"], r["pat"])
    n = sum(1 for r in financials if r["fy"] in summary)
    print(f"  contract_revenue   == AR net revenue   {hits['contract_revenue']} of {n}")
    print(f"  revenue_operations == AR net revenue   {hits['revenue_operations']} of {n}"
          "   <-- expected to fail: revenue_operations is contract x 1.03")
    print(f"  pat                == AR profit        {hits['pat']} of {n}")

    print("\ncross-check 2: statement vs trial-balance workbook (raw rupees)")
    tb = trial_balance_by_fy()
    if not tb:
        print("  trial-balance workbook unavailable   <-- INCOMPLETE")
    else:
        pnl = bs = 0
        n_tb = 0
        for r in financials:
            row = tb.get(r["fy"])
            if not row:
                continue
            n_tb += 1
            pnl += (_rounds_to_lakh(row["revenue"], r["contract_revenue"])
                    and _rounds_to_lakh(row["profit"], r["pat"]))
            recv = row["balances"].get("Accounts Receivable")
            bs += _rounds_to_lakh(abs(recv) if recv is not None else None,
                                  r["trade_receivables"])
        print(f"  P&L (contract revenue AND profit)      {pnl} of {n_tb}")
        print(f"  balance sheet (trade receivables)      {bs} of {n_tb}"
              "   <-- FINDING: the statements' balance-sheet extract is"
              " anchored to no other source")

    print("\ncross-check 3: comparative column vs the prior statement")
    matched = 0
    pairs = 0
    for prev, cur in zip(financials, financials[1:]):
        pairs += 1
        matched += cur.get("contract_revenue_previous") == prev.get("contract_revenue")
    print(f"  previous-year column == prior FY current column   {matched} of {pairs}"
          "   <-- FINDING: the comparative column is a fabricated 0.85x of the")
    print("      current column, NOT the prior year. Do not read *_previous as a")
    print("      fact about the prior FY, and do not treat DOC-FS-2019's")
    print("      comparative as coverage of FY2018-19.")

    print("\ncross-check 4: printed PBT vs the header's own formula (A - B)")
    agree = sum(1 for r in financials
                if r["pbt"] == (r["revenue_operations"] or 0) - (r["total_expenses"] or 0))
    print(f"  pbt == revenue_operations - total_expenses        {agree} of {len(financials)}"
          "   <-- FINDING: recorded as printed")
    consistent = sum(1 for r in financials if r["pbt"] - r["tax"] == r["pat"])
    print(f"  pbt - tax == pat (the printed triple)             {consistent} of {len(financials)}")

    print("\ncross-check 5: compliance oracle DOC-CM-0005 "
          "('minimum average turnover Rs. 150 Cr, FY 2019-2025')")
    turnovers = [r["contract_revenue"] for r in financials
                 if r["contract_revenue"] is not None]
    if turnovers:
        avg = sum(turnovers) / len(turnovers)
        print(f"  mean contract_revenue over {len(turnovers)} FYs = Rs "
              f"{money.format_money(int(avg))} = {avg / 1e7:.2f} Cr  -> "
              f"{'CLEARS' if avg >= 150e7 else 'FAILS'} the 150 Cr gate "
              f"(matrix says Complied)")

    print("\nannual reports")
    for r in reports:
        print(f"  {r['source_doc']}  FY{r['fy']}  gross "
              f"{money.format_money(r['gross_billings'])}  net "
              f"{money.format_money(r['net_revenue'])}  profit "
              f"{money.format_money(r['profit'])}")
    print("=" * 78)


if __name__ == "__main__":
    blocks = build(pdftext.load_corpus())
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(blocks, ensure_ascii=False, indent=2),
                   encoding="utf-8")
    audit(blocks)
    print(f"wrote {sum(len(v) for v in blocks.values())} records -> {OUT}")
